"""Convert the Min-Max Program tracker spreadsheet into src/program.json.

Usage: python3 scripts/parse_xlsx.py Min-Max_Program_4x.xlsx src/program.json
"""
import datetime
import json
import re
import sys

import openpyxl

HEAVY = {"Squat (Your Choice)", "Barbell Incline Press", "Leg Press"}


def rng(v):
    # Excel turned ranges like "6-8" into dates (June 8); turn them back.
    if isinstance(v, datetime.datetime):
        return f"{v.month}-{v.day}"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return None if v is None else str(v).strip()


def link(cell):
    return cell.hyperlink.target if cell.hyperlink else None


def parse(path):
    ws = openpyxl.load_workbook(path).active
    weeks, week, day, block, label = [], None, None, None, None
    for r in range(1, ws.max_row + 1):
        b, c = ws.cell(r, 2).value, ws.cell(r, 3).value
        if isinstance(b, str) and b.startswith("Block"):
            block = int(b.split()[1])
            continue
        if isinstance(b, str) and b.strip() in ("Intro Week", "Deload Week"):
            label = b.strip()
            continue
        if isinstance(b, str) and b.startswith("Week") and c == "Exercise":
            week = {"week": int(b.split()[1]), "block": block, "label": label, "days": []}
            label = None
            weeks.append(week)
            continue
        if week is None or not c or c == "Exercise":
            continue
        if b and "Rest" not in str(b):
            day = {"name": b.strip(), "exercises": []}
            week["days"].append(day)
        tech = rng(ws.cell(r, 4).value)
        day["exercises"].append({
            "name": str(c).strip(),
            "link": link(ws.cell(r, 3)),
            "tech": None if tech == "N/A" else tech,
            "warmup": rng(ws.cell(r, 5).value),
            "sets": int(ws.cell(r, 6).value or 0),
            "reps": rng(ws.cell(r, 7).value),
            "rir": [rng(ws.cell(r, 12).value), rng(ws.cell(r, 13).value)],
            "rest": rng(ws.cell(r, 14).value),
            "subs": [{"name": rng(ws.cell(r, k).value), "link": link(ws.cell(r, k))}
                     for k in (15, 16) if rng(ws.cell(r, k).value) not in (None, "N/A", "See Notes")],
            "notes": rng(ws.cell(r, 17).value),
            "logged": [[ws.cell(r, 8).value, ws.cell(r, 9).value], [ws.cell(r, 10).value, ws.cell(r, 11).value]],
        })
    return weeks


def compact(weeks):
    # Only technique and RIR change week to week, so store the exercise once per day.
    days = []
    for di, d in enumerate(weeks[0]["days"]):
        exs = []
        for ei, e in enumerate(d["exercises"]):
            per_week = [w["days"][di]["exercises"][ei] for w in weeks]
            ex = {k: e[k] for k in ("name", "link", "warmup", "sets", "reps", "rest", "subs", "notes")}
            ex["rir"] = [[x for x in p["rir"][: e["sets"]]] for p in per_week]
            ex["tech"] = [p["tech"] for p in per_week]
            if e["name"] in HEAVY:
                ex["heavy"] = True
            if e["name"].startswith("Dead Hang"):
                ex["timed"] = True
                ex["optional"] = True
            m = re.match(r"This can be an? (.*)\.$", e["notes"] or "")
            if m and not ex["subs"]:
                # "Your choice" lifts list their options in the notes instead.
                ex["subs"] = [{"name": n.strip(), "link": None}
                              for n in re.split(r", (?:or )?", m.group(1))]
            exs.append(ex)
        days.append({"name": d["name"], "exercises": exs})
    return {
        "weeks": [{"week": w["week"], "block": w["block"], "label": w["label"]} for w in weeks],
        "days": days,
    }


if __name__ == "__main__":
    data = compact(parse(sys.argv[1]))
    with open(sys.argv[2], "w") as f:
        json.dump(data, f, indent=1, ensure_ascii=False)
